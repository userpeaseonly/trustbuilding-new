import docx

doc = docx.Document('/app/temp_template.docx')

# 1. Replace __%
for p in doc.paragraphs:
    if "__%" in p.text:
        new_text = p.text.replace("__%", "«Бошланғич_тўлов_фоизи»%", 1)
        new_text = new_text.replace("__%", "«Қолган_тўлов_фоизи»%", 1)
        p.text = new_text

# 2. Fix the payment table
payment_table = None
for table in doc.tables:
    if len(table.rows) > 0 and any('тўлов номи' in c.text.lower() for c in table.rows[0].cells):
        payment_table = table
        break

if payment_table:
    # Remove rows 1 to 60 (inclusive)
    # The rows are: 0=header, 1=Бунак, 2..60=1-59 тулов, 61=Жами, 62=Шартнома 2.6, 63=Кадастр
    for i in range(60, 0, -1):
        row = payment_table.rows[i]
        row._element.getparent().remove(row._element)
    
    # Insert our dynamic row
    # We will use docxtpl syntax. Because CustomDocxTemplate ONLY converts predefined tags like «Шартнома_рақами»,
    # we can just use direct jinja2 `{{ s.name }}` for the loop variables.
    # Wait, CustomDocxTemplate uses `jinja_env` and parses the whole document. So `{% tr %}` works!
    
    new_row = payment_table.add_row()
    # Move it to index 1 (right after header)
    payment_table.rows[0]._element.addnext(new_row._element)
    
    new_row.cells[0].text = "{% tr for s in schedule %}"
    new_row.cells[1].text = "{{ s.name }}"
    new_row.cells[2].text = "{{ s.amount }}"
    new_row.cells[3].text = "{{ s.date }}\n{% tr endfor %}"

doc.save('/app/new_template.docx')

