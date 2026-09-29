import docx
from docx.enum.text import WD_LINE_SPACING
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.shared import Pt

def fix_tables(doc_path, output_path):
    doc = docx.Document(doc_path)
    
    for table in doc.tables:
        for row in table.rows:
            # Clear explicit row height if any
            row.height = None
            row.height_rule = None
            
            for cell in row.cells:
                # Vertically center text in cell
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                
                for p in cell.paragraphs:
                    # Remove spacing before and after
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    
                    # Set line spacing to single
                    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    p.paragraph_format.line_spacing = 1.0

    doc.save(output_path)

fix_tables("/tmp/bunyodkorshartnomashablon.docx", "/tmp/bunyodkorshartnomashablon_tables_fixed.docx")
print("Done")
