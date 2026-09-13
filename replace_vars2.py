from docx import Document

doc = Document('docs/shartnoma_latin.docx')

for p in doc.paragraphs:
    for run in p.runs:
        if 'Арсланов Ахрор Жамолович' in run.text:
            run.text = run.text.replace('Арсланов Ахрор Жамолович', '«Kompaniya_rahbari»')
        if 'А.Ж.Арсланов' in run.text:
            run.text = run.text.replace('А.Ж.Арсланов', '«Kompaniya_rahbari_qisqa»')
                
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    if 'Арсланов Ахрор Жамолович' in run.text:
                        run.text = run.text.replace('Арсланов Ахрор Жамолович', '«Kompaniya_rahbari»')
                    if 'А.Ж.Арсланов' in run.text:
                        run.text = run.text.replace('А.Ж.Арсланов', '«Kompaniya_rahbari_qisqa»')

doc.save('docs/shartnoma_latin_v2.docx')
doc.save('contract/templates/contract/ShartnomaTemplate.docx')
print("Done")
