import docx
import re

doc = docx.Document("/home/dragonfire/Documents/for-antigravity/bunyodkorshartnomashablon.docx")

tags = set()
pattern = re.compile(r'«(.*?)»')

def extract_tags(text):
    for match in pattern.findall(text):
        tags.add(match)

for p in doc.paragraphs:
    extract_tags(p.text)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                extract_tags(p.text)

print("Found tags in DOCX:")
for t in tags:
    print(t)
