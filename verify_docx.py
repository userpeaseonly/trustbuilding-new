import docx
doc = docx.Document('/app/new_template.docx')
for p in doc.paragraphs:
    if "фоизи" in p.text:
        print("PARA:", p.text)

for table in doc.tables:
    if len(table.rows) > 0 and any('тўлов номи' in c.text.lower() for c in table.rows[0].cells):
        for row in table.rows:
            print("ROW:", [c.text.strip().replace('\n', ' ') for c in row.cells])

