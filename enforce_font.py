import docx
from docx.shared import Pt
from docx.oxml.ns import qn

def enforce_font(doc_path, output_path):
    doc = docx.Document(doc_path)
    
    def format_runs(paragraphs):
        for p in paragraphs:
            for run in p.runs:
                if not run.font:
                    continue
                run.font.name = 'Cambria'
                if hasattr(run._element, 'rPr') and hasattr(run._element.rPr, 'rFonts'):
                    run._element.rPr.rFonts.set(qn('w:ascii'), 'Cambria')
                    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Cambria')
                    run._element.rPr.rFonts.set(qn('w:cs'), 'Cambria')
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Cambria')
                
                run.font.size = Pt(12)

    format_runs(doc.paragraphs)
    
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                format_runs(cell.paragraphs)
                
    doc.save(output_path)

enforce_font("/tmp/verify_tags2.docx", "/tmp/bunyodkorshartnomashablon_cambria.docx")
print("Done")
