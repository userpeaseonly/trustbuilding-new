import re

with open('contract/views.py', 'r') as f:
    content = f.read()

new_views = """
@login_required
def download_import_template(request):
    \"\"\"Generate a blank Excel template for bulk payment imports.\"\"\"
    if not getattr(request.user, 'is_company', False) and not getattr(request.user, 'is_staff_member', False):
        return redirect('dashboard:home')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Payments Template"

    headers = ['Sana (DD.MM.YYYY)', 'Summa', "To'lov turi"]
    ws.append(headers)

    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = openpyxl.styles.Font(bold=True)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 20

    # Data validation for Payment Type
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1='"Cash,Card,Bank Transfer,Material"', allow_blank=True)
    ws.add_data_validation(dv)
    # Apply to C2:C1000
    dv.add('C2:C1000')

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Payment_Template.xlsx"'
    wb.save(response)
    return response


@login_required
def import_payments_excel(request, pk):
    \"\"\"Process the uploaded Excel template and bulk create payment logs.\"\"\"
    company = get_user_company(request.user)
    contract = get_object_or_404(Contract, pk=pk, company=company)

    if request.method == 'POST':
        if contract.payment_logs.count() > 0:
            messages.error(request, _("Cannot import payments: Contract already has existing payments."))
            return redirect('contract:detail', pk=contract.pk)

        excel_file = request.FILES.get('file')
        if not excel_file or not excel_file.name.endswith('.xlsx'):
            messages.error(request, _("Please upload a valid .xlsx file."))
            return redirect('contract:detail', pk=contract.pk)

        try:
            wb = openpyxl.load_workbook(excel_file)
            ws = wb.active

            type_mapping = {
                'cash': PaymentLog.PAYMENT_TYPE_CASH,
                'card': PaymentLog.PAYMENT_TYPE_CARD,
                'bank transfer': PaymentLog.PAYMENT_TYPE_BANK,
                'material': PaymentLog.PAYMENT_TYPE_MATERIAL,
            }

            import datetime

            payments_created = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                raw_date, raw_amount, raw_type = row[0], row[1], row[2]
                
                # Skip completely empty rows
                if not raw_date and not raw_amount:
                    continue

                try:
                    amount = Decimal(str(raw_amount))
                except:
                    continue  # Skip invalid amount
                
                # Parse Date
                date_paid = timezone.now()
                if isinstance(raw_date, datetime.datetime):
                    date_paid = timezone.make_aware(raw_date) if timezone.is_naive(raw_date) else raw_date
                elif isinstance(raw_date, str):
                    try:
                        parsed = datetime.datetime.strptime(raw_date, '%d.%m.%Y')
                        date_paid = timezone.make_aware(parsed)
                    except:
                        pass

                # Parse Payment Type
                payment_type = PaymentLog.PAYMENT_TYPE_CASH
                if raw_type and str(raw_type).strip().lower() in type_mapping:
                    payment_type = type_mapping[str(raw_type).strip().lower()]

                payment = PaymentLog(
                    contract=contract,
                    amount=amount,
                    payment_type=payment_type,
                    date_paid=date_paid
                )
                payment.save()
                payments_created += 1

            messages.success(request, _(f"Successfully imported {payments_created} payments!"))
        except Exception as e:
            import logging
            logging.getLogger(__name__).error(f"Error importing payments: {e}")
            messages.error(request, _("Failed to process Excel file. Please ensure it matches the template."))

    return redirect('contract:detail', pk=contract.pk)
"""

parts = content.split('@login_required\ndef download_payments_excel')
if len(parts) == 2:
    new_content = parts[0] + new_views + '\n@login_required\ndef download_payments_excel' + parts[1]
    with open('contract/views.py', 'w') as f:
        f.write(new_content)
    print("Success")
else:
    print("Could not find anchor")
