from docx import Document

replace_map = {
    'Шартнома_рақами': 'Shartnoma_raqami',
    'Шартнома_бўлган_сана': 'Shartnoma_kuni',
    'Шартнома_ойи': 'Shartnoma_oyi',
    'Шартнома_бўлган_йил': 'Shartnoma_yili',
    'Сотиб_олувчи_ФИО': 'Xaridor_FIO',
    'Шартнома_суммаси': 'Shartnoma_summasi',
    'Шартнома_суммаси_сўз_билан': 'Shartnoma_summasi_soz_bilan',
    'Блок_': 'Blok',
    'Подъезд': 'Podyezd',
    'Этаж': 'Qavat',
    'Квартира_рақами': 'Kvartira_raqami',
    'Яшаш_хоналар_сони': 'Xonalar_soni',
    'Умумий_майдони': 'Umumiy_maydon',
    'M_1_кв_метр_нархи': 'Bir_kv_metr_narxi',
    'M_1_кв_нархи_сўз_билан': 'Bir_kv_metr_narxi_soz_bilan',
    'Бошланғич_тўлов': 'Boshlangich_tolov',
    'График_бўйича_ойлик_тўлов': 'Oylik_tolov',
    'Сотиб_олувчи_паспорт_серияси': 'Xaridor_pasporti',
    'Паспорти_берилган_вақти': 'Pasport_berilgan_sana',
    'Паспорт_берилган_жойи': 'Pasport_berilgan_joy',
    'Телефон_рақами': 'Telefon_raqami',
    'Сотиб_олувчи_ФИО_қисқартмаси': 'Xaridor_FIO_qisqa'
}

def replace_text(doc):
    for p in doc.paragraphs:
        for run in p.runs:
            for old, new in replace_map.items():
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        for old, new in replace_map.items():
                            if old in run.text:
                                run.text = run.text.replace(old, new)

doc = Document('docs/№314.docx')
replace_text(doc)
doc.save('shartnoma_latin.docx')
print("Saved successfully!")
