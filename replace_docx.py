import docx

def replace_text_in_doc(doc_path, output_path, replacements):
    doc = docx.Document(doc_path)
    
    def replace_in_paragraphs(paragraphs):
        for p in paragraphs:
            for old_text, new_text in replacements.items():
                if old_text in p.text:
                    # Replace across the entire paragraph text
                    # We will just re-assign the text. It might lose inline bolding, but usually acceptable.
                    p.text = p.text.replace(old_text, new_text)

    replace_in_paragraphs(doc.paragraphs)
    
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                replace_in_paragraphs(cell.paragraphs)
                
    doc.save(output_path)

replacements = {
    "«BUNYODKOR SAMARQAND»": "«Пудратчи»",
    "BUNYODKOR SAMARQAND": "«Пудратчи»",
    "Абдуллаев Фарходбек Шухратбекович": "«Пудратчи_раҳбари»",
    "НАРПАЙ Т., «АГРОБАНК» АТБ": "«Пудратчи_банки»",
    "МФО: 00267": "МФО: «Пудратчи_МФО»",
    "х/р: 20208000600550335001": "х/р: «Пудратчи_ХР»",
    "ИНН: 304561089": "ИНН: «Пудратчи_ИНН»",
    "Самарқанд вилояти Нарпай тумани Нарпай МФЙ Ойтамғали кўчаси 127-уй": "«Пудратчи_Манзили»",
    "14:06:41:02:02:1797": "«Кадастр_рақами»",
    "12 620": "«Ер_майдони»",
    "106.68": "«Қурилиш_ости_майдони»",
    "««Шартнома_бўлган_сана»»": "«Шартнома_бўлган_сана»"
}

replace_text_in_doc("/tmp/bunyodkorshartnomashablon.docx", "/tmp/bunyodkorshartnomashablon_updated.docx", replacements)
print("Done")
