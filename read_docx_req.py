import docx
doc_path = '/app/temp_template.docx'
doc = docx.Document(doc_path)
for i in range(max(0, len(doc.paragraphs)-30), len(doc.paragraphs)):
    if doc.paragraphs[i].text.strip():
        print(f"[{i}] {doc.paragraphs[i].text}")

for table in doc.tables:
    for row in table.rows:
        row_text = " | ".join([cell.text.replace("\n", " ") for cell in row.cells])
        print("TABLE_ROW:", row_text)

