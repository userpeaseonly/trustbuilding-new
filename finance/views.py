import calendar
from datetime import datetime
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db.models import Sum, Q, Count
from django.http import HttpResponse, Http404, JsonResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse
from django.utils import timezone
from django.utils.translation import gettext as _
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from building.models import Building, Apartment
from building.views import get_user_company
from contract.models import Contract, PaymentRecord, PaymentLog
from otp.services import send_sms
from otp.transliteration import build_overdue_sms, build_upcoming_reminder_sms
from users.permissions import require_permission
from .models import Wallet, ExpenseCategory, FinancialTransaction


DEFAULT_EXPENSE_CATEGORIES = [
    "Материалы (Бетон, Арматура, Кирпич)",
    "Строительно-монтажные работы (Субподрядчики)",
    "Аренда спецтехники и оборудования",
    "Зарплата и авансы строителям",
    "Маркетинг, реклама и комиссии продаж",
    "Административные и офисные расходы",
    "Проектные работы, кадастр и разрешения",
    "Налоги и банковские комиссии",
    "Прочие расходы",
]


def ensure_company_defaults(company):
    """Ensure essential cash desks and expense categories exist for the company."""
    if not Wallet.objects.filter(company=company).exists():
        Wallet.objects.create(
            company=company,
            name=_("Главная касса (Наличные)"),
            wallet_type='CASH',
            currency='UZS',
            is_default=True
        )
        Wallet.objects.create(
            company=company,
            name=_("Расчетный счет (Банк)"),
            wallet_type='BANK',
            currency='UZS',
            is_default=False
        )

    for cat_name in DEFAULT_EXPENSE_CATEGORIES:
        ExpenseCategory.objects.get_or_create(
            company=company,
            name=cat_name,
            category_type='EXPENSE',
            defaults={'is_system': True}
        )
        
    ExpenseCategory.objects.get_or_create(
        company=company,
        name=_("Оплата по договорам (Поступления от клиентов)"),
        category_type='INCOME',
        defaults={'is_system': True}
    )


@login_required
def finance_overview(request):
    """
    Main Cash Flow / ДДС Hub:
    - Wallet balances
    - Inflows vs Outflows
    - Ledger table with filtering & Excel export
    """
    company = get_user_company(request.user)
    if not company:
        messages.error(request, _("Company profile not found."))
        return redirect('dashboard:home')
        
    ensure_company_defaults(company)
    
    # Active tab: 'cashflow', 'debtors', 'costing'
    active_tab = request.GET.get('tab', 'cashflow')
    
    # Date filters
    today = timezone.localdate()
    start_date_str = request.GET.get('start_date')
    end_date_str = request.GET.get('end_date')
    
    if start_date_str:
        try:
            start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date()
        except ValueError:
            start_date = today.replace(day=1)
    else:
        # Default to 1st day of current month
        start_date = today.replace(day=1)

    if end_date_str:
        try:
            end_date = datetime.strptime(end_date_str, '%Y-%m-%d').date()
        except ValueError:
            end_date = today
    else:
        end_date = today

    # 1. Wallets Summary
    wallets = Wallet.objects.filter(company=company, is_active=True)
    wallet_balances = []
    total_cash = 0
    total_bank = 0
    total_liquid = 0

    for w in wallets:
        bal = w.current_balance
        wallet_balances.append({
            'wallet': w,
            'balance': bal,
        })
        if w.currency == 'UZS':
            if w.wallet_type == 'CASH':
                total_cash += bal
            elif w.wallet_type in ['BANK', 'CARD']:
                total_bank += bal
            total_liquid += bal

    # 2. Period Inflows & Outflows
    period_txs = FinancialTransaction.objects.filter(
        company=company,
        date__range=[start_date, end_date]
    )
    period_income = period_txs.filter(transaction_type='INCOME').aggregate(s=Sum('amount'))['s'] or 0
    period_expense = period_txs.filter(transaction_type='EXPENSE').aggregate(s=Sum('amount'))['s'] or 0
    net_cash_flow = period_income - period_expense

    # 3. Filtered Ledger
    ledger_qs = FinancialTransaction.objects.filter(
        company=company,
        date__range=[start_date, end_date]
    ).select_related('wallet', 'target_wallet', 'category', 'building', 'contract', 'created_by')

    filter_type = request.GET.get('tx_type', '')
    if filter_type in ['INCOME', 'EXPENSE', 'TRANSFER']:
        ledger_qs = ledger_qs.filter(transaction_type=filter_type)

    filter_wallet = request.GET.get('wallet_id', '')
    if filter_wallet:
        ledger_qs = ledger_qs.filter(Q(wallet_id=filter_wallet) | Q(target_wallet_id=filter_wallet))

    filter_category = request.GET.get('category_id', '')
    if filter_category:
        ledger_qs = ledger_qs.filter(category_id=filter_category)

    filter_building = request.GET.get('building_id', '')
    if filter_building:
        ledger_qs = ledger_qs.filter(building_id=filter_building)

    q_search = request.GET.get('q', '').strip()
    if q_search:
        ledger_qs = ledger_qs.filter(
            Q(description__icontains=q_search) |
            Q(contract__contract_id__icontains=q_search) |
            Q(contract__customer__full_name__icontains=q_search) |
            Q(contract__customer__phone_number__icontains=q_search)
        )

    # Excel Export
    if request.GET.get('export') == 'excel':
        return export_ledger_excel(ledger_qs, start_date, end_date)

    paginator = Paginator(ledger_qs, 25)
    page_number = request.GET.get('page', 1)
    transactions_page = paginator.get_page(page_number)

    # Dropdowns for Filters & Modals
    categories = ExpenseCategory.objects.filter(company=company)
    buildings = Building.objects.filter(company=company)

    # Tab 2 data: Debtors & Inflow Forecast
    debtors_data = []
    forecast_months = []
    month_planned_inflow = 0
    month_actual_inflow = 0
    collection_rate = 0
    aging_stats = {
        'all': {'count': 0, 'amount': 0},
        '1_15': {'count': 0, 'amount': 0},
        '16_30': {'count': 0, 'amount': 0},
        '31_60': {'count': 0, 'amount': 0},
        '60_plus': {'count': 0, 'amount': 0},
    }
    selected_bucket = request.GET.get('bucket', 'all')
    debtor_building_id = request.GET.get('debtor_building', '')

    if active_tab == 'debtors':
        # 1. 6-Month Rolling Forecast
        month_names_uz = {
            1: "Yanvar", 2: "Fevral", 3: "Mart", 4: "Aprel", 5: "May", 6: "Iyun",
            7: "Iyul", 8: "Avgust", 9: "Sentabr", 10: "Oktabr", 11: "Noyabr", 12: "Dekabr"
        }
        month_names_ru = {
            1: "Январь", 2: "Февраль", 3: "Март", 4: "Апрель", 5: "Май", 6: "Июнь",
            7: "Июль", 8: "Август", 9: "Сентябрь", 10: "Октябрь", 11: "Ноябрь", 12: "Декабрь"
        }

        for i in range(6):
            target_month = ((today.month - 1 + i) % 12) + 1
            target_year = today.year + ((today.month - 1 + i) // 12)

            m_records = PaymentRecord.objects.filter(
                contract__company=company,
                contract__status='ACTIVE',
                due_date__year=target_year,
                due_date__month=target_month
            )
            m_planned = m_records.aggregate(s=Sum('plan_amount'))['s'] or 0
            m_paid = m_records.aggregate(s=Sum('paid_amount'))['s'] or 0
            m_contracts_count = m_records.values('contract').distinct().count()
            m_rate = round((m_paid / m_planned * 100), 1) if m_planned > 0 else 0

            forecast_months.append({
                'year': target_year,
                'month': target_month,
                'month_name': calendar.month_name[target_month],
                'month_name_ru': month_names_ru.get(target_month, ''),
                'month_name_uz': month_names_uz.get(target_month, ''),
                'planned': m_planned,
                'paid': m_paid,
                'contracts_count': m_contracts_count,
                'is_current': (i == 0),
                'rate': m_rate,
            })

        if forecast_months:
            month_planned_inflow = forecast_months[0]['planned']
            month_actual_inflow = forecast_months[0]['paid']
            collection_rate = forecast_months[0]['rate']

        # 2. Overdue records & Aging Matrix
        all_overdue = PaymentRecord.objects.filter(
            contract__company=company,
            contract__status='ACTIVE',
            due_date__lt=today,
            debt__gt=0
        ).select_related(
            'contract', 'contract__customer',
            'contract__apartment', 'contract__apartment__building'
        ).order_by('due_date')

        filtered_debtors = []
        for r in all_overdue:
            days_late = (today - r.due_date).days
            r.days_late = days_late
            debt_amt = r.debt

            aging_stats['all']['count'] += 1
            aging_stats['all']['amount'] += debt_amt

            if 1 <= days_late <= 15:
                bucket_key = '1_15'
                r.bucket_label = '1-15'
                r.badge_class = 'bg-amber-50 text-amber-700 border-amber-200'
            elif 16 <= days_late <= 30:
                bucket_key = '16_30'
                r.bucket_label = '16-30'
                r.badge_class = 'bg-orange-50 text-orange-700 border-orange-200'
            elif 31 <= days_late <= 60:
                bucket_key = '31_60'
                r.bucket_label = '31-60'
                r.badge_class = 'bg-rose-50 text-rose-700 border-rose-200'
            else:
                bucket_key = '60_plus'
                r.bucket_label = '60+'
                r.badge_class = 'bg-purple-50 text-purple-700 border-purple-200'

            aging_stats[bucket_key]['count'] += 1
            aging_stats[bucket_key]['amount'] += debt_amt

            # Apply bucket filter
            if selected_bucket != 'all' and selected_bucket != bucket_key:
                continue

            # Apply building filter
            if debtor_building_id:
                apt = getattr(r.contract, 'apartment', None)
                if not apt or str(apt.building_id) != str(debtor_building_id):
                    continue

            filtered_debtors.append(r)

        debtors_data = filtered_debtors[:100]

    # Tab 3 data: Building Costing & Unit Economics (Phase 3)
    buildings_costing = []
    portfolio_costing = {
        'total_area': 0.0,
        'sold_area': 0.0,
        'available_area': 0.0,
        'sold_pct': 0.0,
        'total_apts': 0,
        'sold_apts': 0,
        'available_apts': 0,
        'contracted_revenue': 0.0,
        'collected_revenue': 0.0,
        'pending_receivables': 0.0,
        'total_expenses': 0.0,
        'construction_expenses': 0.0,
        'overhead_expenses': 0.0,
        'net_cash_margin': 0.0,
        'projected_profit': 0.0,
        'avg_sale_price_sqm': 0.0,
        'avg_cost_sqm': 0.0,
        'margin_per_sqm': 0.0,
        'profit_margin_pct': 0.0,
    }
    selected_costing_bld_id = request.GET.get('costing_bld', '')
    selected_costing_bld = None
    category_breakdown = []
    recent_costing_txs = []
    target_costing_area = 0.0
    target_total_expense = 0.0

    if active_tab == 'costing':
        # 1. Apartments pre-aggregation
        apts_qs = Apartment.objects.filter(building__company=company, is_real=True)
        bld_apt_stats = {}
        for a in apts_qs.values('building_id').annotate(
            agg_total_area=Sum('total_area'),
            agg_total_apts=Count('id'),
            agg_sold_area=Sum('total_area', filter=Q(status='SOLD')),
            agg_sold_apts=Count('id', filter=Q(status='SOLD')),
        ):
            bld_apt_stats[a['building_id']] = a

        # 2. Contracts pre-aggregation
        contract_stats = {}
        for c in Contract.objects.filter(company=company, status='ACTIVE').values('apartment__building_id').annotate(
            agg_contracted_rev=Sum('total_amount'),
            agg_contracts_count=Count('id')
        ):
            contract_stats[c['apartment__building_id']] = c

        # 3. Financial Transactions pre-aggregation
        tx_stats = {}
        for t in FinancialTransaction.objects.filter(company=company).values('building_id', 'transaction_type').annotate(
            agg_amount=Sum('amount')
        ):
            b_id = t['building_id']
            t_type = t['transaction_type']
            if b_id not in tx_stats:
                tx_stats[b_id] = {'INCOME': 0.0, 'EXPENSE': 0.0}
            tx_stats[b_id][t_type] = float(t['agg_amount'] or 0)

        overhead_expenses = tx_stats.get(None, {}).get('EXPENSE', 0.0)

        # 4. Building-by-building metrics
        all_buildings = buildings.order_by('name', 'block_number')
        for bld in all_buildings:
            apt_st = bld_apt_stats.get(bld.id, {})
            c_st = contract_stats.get(bld.id, {})
            t_st = tx_stats.get(bld.id, {'INCOME': 0.0, 'EXPENSE': 0.0})

            b_tot_area = float(apt_st.get('agg_total_area') or 0.0)
            b_sold_area = float(apt_st.get('agg_sold_area') or 0.0)
            b_avail_area = max(0.0, b_tot_area - b_sold_area)
            b_sold_pct = round((b_sold_area / b_tot_area * 100), 1) if b_tot_area > 0 else 0.0

            b_tot_apts = int(apt_st.get('agg_total_apts') or 0)
            b_sold_apts = int(apt_st.get('agg_sold_apts') or 0)
            b_avail_apts = b_tot_apts - b_sold_apts

            b_contracted = float(c_st.get('agg_contracted_rev') or 0.0)
            b_collected = float(t_st.get('INCOME', 0.0))
            b_pending = max(0.0, b_contracted - b_collected)
            b_expense = float(t_st.get('EXPENSE', 0.0))

            b_avg_sale_sqm = round(b_contracted / b_sold_area, 0) if b_sold_area > 0 else 0.0
            b_cost_sqm = round(b_expense / b_tot_area, 0) if b_tot_area > 0 else 0.0
            b_margin_sqm = b_avg_sale_sqm - b_cost_sqm
            b_profit_pct = round((b_margin_sqm / b_avg_sale_sqm * 100), 1) if b_avg_sale_sqm > 0 else 0.0
            b_net_margin = b_collected - b_expense
            b_projected_profit = b_contracted - b_expense

            is_sel = (selected_costing_bld_id == str(bld.id))
            if is_sel:
                selected_costing_bld = bld

            b_data = {
                'building': bld,
                'total_area': b_tot_area,
                'sold_area': b_sold_area,
                'available_area': b_avail_area,
                'sold_pct': b_sold_pct,
                'total_apts': b_tot_apts,
                'sold_apts': b_sold_apts,
                'available_apts': b_avail_apts,
                'contracted_rev': b_contracted,
                'collected_rev': b_collected,
                'pending_receivables': b_pending,
                'actual_expenses': b_expense,
                'cost_per_sqm': b_cost_sqm,
                'avg_sale_price_sqm': b_avg_sale_sqm,
                'margin_per_sqm': b_margin_sqm,
                'profit_margin_pct': b_profit_pct,
                'net_cash_margin': b_net_margin,
                'projected_profit': b_projected_profit,
                'is_selected': is_sel,
            }
            buildings_costing.append(b_data)

            # Portfolio accumulations
            portfolio_costing['total_area'] += b_tot_area
            portfolio_costing['sold_area'] += b_sold_area
            portfolio_costing['total_apts'] += b_tot_apts
            portfolio_costing['sold_apts'] += b_sold_apts
            portfolio_costing['contracted_revenue'] += b_contracted
            portfolio_costing['collected_revenue'] += b_collected
            portfolio_costing['construction_expenses'] += b_expense

        # Portfolio derived metrics
        p_tot_area = portfolio_costing['total_area']
        p_sold_area = portfolio_costing['sold_area']
        portfolio_costing['available_area'] = max(0.0, p_tot_area - p_sold_area)
        portfolio_costing['sold_pct'] = round((p_sold_area / p_tot_area * 100), 1) if p_tot_area > 0 else 0.0
        portfolio_costing['available_apts'] = portfolio_costing['total_apts'] - portfolio_costing['sold_apts']
        portfolio_costing['pending_receivables'] = max(0.0, portfolio_costing['contracted_revenue'] - portfolio_costing['collected_revenue'])
        portfolio_costing['overhead_expenses'] = overhead_expenses
        portfolio_costing['total_expenses'] = portfolio_costing['construction_expenses'] + overhead_expenses
        portfolio_costing['net_cash_margin'] = portfolio_costing['collected_revenue'] - portfolio_costing['total_expenses']
        portfolio_costing['projected_profit'] = portfolio_costing['contracted_revenue'] - portfolio_costing['total_expenses']

        p_con_rev = portfolio_costing['contracted_revenue']
        portfolio_costing['avg_sale_price_sqm'] = round(p_con_rev / p_sold_area, 0) if p_sold_area > 0 else 0.0
        portfolio_costing['avg_cost_sqm'] = round(portfolio_costing['total_expenses'] / p_tot_area, 0) if p_tot_area > 0 else 0.0
        portfolio_costing['margin_per_sqm'] = portfolio_costing['avg_sale_price_sqm'] - portfolio_costing['avg_cost_sqm']
        p_avg_sale = portfolio_costing['avg_sale_price_sqm']
        portfolio_costing['profit_margin_pct'] = round((portfolio_costing['margin_per_sqm'] / p_avg_sale * 100), 1) if p_avg_sale > 0 else 0.0

        # 5. Category Breakdown & Recent Expenses
        costing_tx_qs = FinancialTransaction.objects.filter(
            company=company,
            transaction_type='EXPENSE'
        )
        if selected_costing_bld:
            costing_tx_qs = costing_tx_qs.filter(building=selected_costing_bld)
            target_b_data = next((x for x in buildings_costing if x['building'].id == selected_costing_bld.id), None)
            target_costing_area = target_b_data['total_area'] if target_b_data else 0.0
        else:
            target_costing_area = p_tot_area

        target_total_expense = float(costing_tx_qs.aggregate(s=Sum('amount'))['s'] or 0.0)

        cat_agg = costing_tx_qs.values('category__id', 'category__name').annotate(
            total=Sum('amount'),
            tx_count=Count('id')
        ).order_by('-total')

        for c in cat_agg:
            c_amt = float(c['total'] or 0.0)
            c_cost_sqm = round(c_amt / target_costing_area, 0) if target_costing_area > 0 else 0.0
            c_share = round((c_amt / target_total_expense * 100), 1) if target_total_expense > 0 else 0.0
            category_breakdown.append({
                'id': c['category__id'],
                'name': c['category__name'] or _("Без категории"),
                'amount': c_amt,
                'cost_per_sqm': c_cost_sqm,
                'share_pct': c_share,
                'tx_count': c['tx_count'],
            })

        recent_costing_txs = costing_tx_qs.select_related(
            'wallet', 'category', 'building', 'created_by'
        ).order_by('-date', '-created_at')[:20]

    context = {
        'active_tab': active_tab,
        'today': today,
        'start_date': start_date,
        'end_date': end_date,
        'wallets': wallets,
        'wallet_balances': wallet_balances,
        'total_cash': total_cash,
        'total_bank': total_bank,
        'total_liquid': total_liquid,
        'period_income': period_income,
        'period_expense': period_expense,
        'net_cash_flow': net_cash_flow,
        'transactions': transactions_page,
        'categories': categories,
        'buildings': buildings,
        'filter_type': filter_type,
        'filter_wallet': filter_wallet,
        'filter_category': filter_category,
        'filter_building': filter_building,
        'q_search': q_search,
        # Tab 2: Forecast & Aging
        'forecast_months': forecast_months,
        'month_planned_inflow': month_planned_inflow,
        'month_actual_inflow': month_actual_inflow,
        'collection_rate': collection_rate,
        'debtors_data': debtors_data,
        'aging_stats': aging_stats,
        'selected_bucket': selected_bucket,
        'debtor_building_id': debtor_building_id,
        # Tab 3: Costing & Unit Economics
        'buildings_costing': buildings_costing,
        'portfolio_costing': portfolio_costing,
        'selected_costing_bld_id': selected_costing_bld_id,
        'selected_costing_bld': selected_costing_bld,
        'category_breakdown': category_breakdown,
        'target_costing_area': target_costing_area,
        'target_total_expense': target_total_expense,
        'recent_costing_txs': recent_costing_txs,
    }
    return render(request, 'finance/overview.html', context)


@login_required
def send_debtor_reminder(request, pk):
    """Trigger an instant SMS reminder to an overdue debtor."""
    if request.method != 'POST':
        return redirect(f"{reverse('finance:overview')}?tab=debtors")

    company = get_user_company(request.user)
    if not company:
        messages.error(request, _("Company profile not found."))
        return redirect('finance:overview')

    record = get_object_or_404(
        PaymentRecord,
        id=pk,
        contract__company=company,
        contract__status='ACTIVE'
    )

    if record.debt <= 0:
        messages.warning(request, _("This record has no outstanding debt."))
        return redirect(f"{reverse('finance:overview')}?tab=debtors")

    customer = record.contract.customer
    customer_phone = str(customer.phone_number)
    customer_name = customer.full_name or _("Customer")
    days_late = (timezone.localdate() - record.due_date).days

    if days_late > 0:
        message = build_overdue_sms(
            customer_full_name=customer.full_name,
            contract_id=record.contract.id,
            due_date=record.due_date,
            debt_amount=record.debt
        )
    else:
        message = build_upcoming_reminder_sms(
            contract_id=record.contract.id,
            due_date=record.due_date,
            amount=record.debt
        )

    try:
        sent = send_sms(customer_phone, message, company=company)
        if sent:
            messages.success(
                request,
                _("SMS reminder sent to %(name)s (%(phone)s) successfully!") % {
                    'name': customer_name,
                    'phone': customer_phone
                }
            )
        else:
            messages.warning(
                request,
                _("Could not deliver SMS to %(name)s (%(phone)s). Gateway returned failure.") % {
                    'name': customer_name,
                    'phone': customer_phone
                }
            )
    except Exception as e:
        messages.error(request, _("Failed to send SMS: %(error)s") % {'error': str(e)})

    return redirect(f"{reverse('finance:overview')}?tab=debtors")


@login_required
def add_transaction(request):
    """Create a manual Income, Expense, or Transfer transaction."""
    if request.method != 'POST':
        return redirect('finance:overview')

    company = get_user_company(request.user)
    tx_type = request.POST.get('transaction_type')
    amount_raw = request.POST.get('amount', '').replace(' ', '').replace(',', '.')
    date_str = request.POST.get('date')
    wallet_id = request.POST.get('wallet')
    description = request.POST.get('description', '').strip()

    try:
        amount = float(amount_raw)
        if amount <= 0:
            raise ValueError()
    except ValueError:
        messages.error(request, _("Please provide a valid positive amount."))
        return redirect('finance:overview')

    try:
        date = datetime.strptime(date_str, '%Y-%m-%d').date()
    except (ValueError, TypeError):
        date = timezone.localdate()

    wallet = get_object_or_404(Wallet, id=wallet_id, company=company)

    if tx_type == 'TRANSFER':
        target_wallet_id = request.POST.get('target_wallet')
        if not target_wallet_id or target_wallet_id == wallet_id:
            messages.error(request, _("Please choose a different destination account for transfer."))
            return redirect('finance:overview')
        target_wallet = get_object_or_404(Wallet, id=target_wallet_id, company=company)
        
        FinancialTransaction.objects.create(
            company=company,
            transaction_type='TRANSFER',
            wallet=wallet,
            target_wallet=target_wallet,
            amount=amount,
            date=date,
            description=description or _("Перевод средств (Инкассация)"),
            created_by=request.user
        )
        messages.success(request, _("Transfer executed successfully!"))
        return redirect('finance:overview')

    elif tx_type in ['INCOME', 'EXPENSE']:
        category_id = request.POST.get('category')
        building_id = request.POST.get('building')
        
        category = ExpenseCategory.objects.filter(id=category_id, company=company).first() if category_id else None
        building = Building.objects.filter(id=building_id, company=company).first() if building_id else None

        FinancialTransaction.objects.create(
            company=company,
            transaction_type=tx_type,
            wallet=wallet,
            amount=amount,
            date=date,
            category=category,
            building=building,
            description=description,
            created_by=request.user
        )
        messages.success(request, _("Operation recorded successfully!"))
        next_tab = request.POST.get('next_tab', '')
        costing_bld = request.POST.get('costing_bld', '') or (building_id if next_tab == 'costing' else '')
        redirect_url = reverse('finance:overview')
        params = []
        if next_tab:
            params.append(f"tab={next_tab}")
        if costing_bld:
            params.append(f"costing_bld={costing_bld}")
        if params:
            redirect_url += "?" + "&".join(params)
        return redirect(redirect_url)

    messages.error(request, _("Invalid operation type."))
    return redirect('finance:overview')


@login_required
def delete_transaction(request, pk):
    """Delete a manual transaction (contract-synced receipts cannot be deleted directly here)."""
    company = get_user_company(request.user)
    tx = get_object_or_404(FinancialTransaction, id=pk, company=company)

    if tx.payment_log:
        messages.error(request, _("This record is linked to a Contract payment receipt. To modify it, please go to the Contract page."))
        return redirect('finance:overview')

    tx.delete()
    messages.success(request, _("Record deleted successfully."))
    return redirect('finance:overview')


@login_required
def manage_wallets(request):
    """Manage company cash desks and bank accounts."""
    company = get_user_company(request.user)
    wallets = Wallet.objects.filter(company=company)

    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        wallet_type = request.POST.get('wallet_type', 'CASH')
        currency = request.POST.get('currency', 'UZS')
        initial_balance_raw = request.POST.get('initial_balance', '0').replace(' ', '').replace(',', '.')
        account_number = request.POST.get('account_number', '').strip()
        bank_name = request.POST.get('bank_name', '').strip()

        try:
            initial_balance = float(initial_balance_raw)
        except ValueError:
            initial_balance = 0.0

        if name:
            Wallet.objects.create(
                company=company,
                name=name,
                wallet_type=wallet_type,
                currency=currency,
                initial_balance=initial_balance,
                account_number=account_number,
                bank_name=bank_name
            )
            messages.success(request, _("New account/cash desk created!"))
            return redirect('finance:manage_wallets')

    return render(request, 'finance/manage_wallets.html', {'wallets': wallets})


@login_required
def manage_categories(request):
    """Manage company expense categories."""
    company = get_user_company(request.user)
    categories = ExpenseCategory.objects.filter(company=company)

    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        cat_type = request.POST.get('category_type', 'EXPENSE')
        if name:
            ExpenseCategory.objects.create(
                company=company,
                name=name,
                category_type=cat_type
            )
            messages.success(request, _("New category created!"))
            return redirect('finance:manage_categories')

    return render(request, 'finance/manage_categories.html', {'categories': categories})


def export_ledger_excel(queryset, start_date, end_date):
    """Generate clean Excel export of financial transactions."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cash Flow Ledger"

    headers = [
        _("Date"),
        _("Type"),
        _("Cash Desk / Account"),
        _("Category"),
        _("Building / Object"),
        _("Amount (UZS)"),
        _("Description"),
        _("Recorded By")
    ]
    ws.append(headers)

    # Style Header
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    row_idx = 2
    for tx in queryset:
        tx_type_str = tx.get_transaction_type_display()
        if tx.transaction_type == 'TRANSFER':
            wallet_str = f"{tx.wallet.name} -> {tx.target_wallet.name if tx.target_wallet else ''}"
        else:
            wallet_str = tx.wallet.name

        cat_str = tx.category.name if tx.category else ""
        bld_str = tx.building.name if tx.building else ""
        rec_by = tx.created_by.full_name or str(tx.created_by.phone_number) if tx.created_by else ""

        amount_val = float(tx.amount) if tx.transaction_type == 'INCOME' else -float(tx.amount)

        ws.append([
            tx.date.strftime('%d.%m.%Y'),
            str(tx_type_str),
            wallet_str,
            cat_str,
            bld_str,
            amount_val,
            tx.description,
            rec_by
        ])

        # Style amount cell color
        amt_cell = ws.cell(row=row_idx, column=6)
        amt_cell.number_format = '#,##0'
        if tx.transaction_type == 'INCOME':
            amt_cell.font = Font(color="16A34A", bold=True)
        elif tx.transaction_type == 'EXPENSE':
            amt_cell.font = Font(color="DC2626", bold=True)

        for col_num in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col_num).border = thin_border

        row_idx += 1

    # Adjust Column Widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"trustbuilding_cashflow_{start_date.strftime('%Y%m%d')}_{end_date.strftime('%Y%m%d')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
