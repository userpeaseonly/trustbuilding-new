import os
import django
from docx import Document

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'application.settings')
django.setup()

from contract.models import ContractTemplate

for tmpl in ContractTemplate.objects.all():
    try:
        if not tmpl.file:
            continue
            
        doc = Document(tmpl.file.path)
        modified = False
        
        for p in doc.paragraphs:
            if 'Арсланов Ахрор Жамолович' in p.text:
                for run in p.runs:
                    if 'Арсланов Ахрор Жамолович' in run.text:
                        run.text = run.text.replace('Арсланов Ахрор Жамолович', '«Kompaniya_rahbari»')
                        modified = True
            if 'А.Ж.Арсланов' in p.text:
                for run in p.runs:
                    if 'А.Ж.Арсланов' in run.text:
                        run.text = run.text.replace('А.Ж.Арсланов', '«Kompaniya_rahbari_qisqa»')
                        modified = True
                        
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if 'Арсланов Ахрор Жамолович' in p.text:
                            for run in p.runs:
                                if 'Арсланов Ахрор Жамолович' in run.text:
                                    run.text = run.text.replace('Арсланов Ахрор Жамолович', '«Kompaniya_rahbari»')
                                    modified = True
                        if 'А.Ж.Арсланов' in p.text:
                            for run in p.runs:
                                if 'А.Ж.Арсланов' in run.text:
                                    run.text = run.text.replace('А.Ж.Арсланов', '«Kompaniya_rahbari_qisqa»')
                                    modified = True
                                    
        if modified:
            doc.save(tmpl.file.path)
            print(f"Fixed director in template {tmpl.id}")
    except Exception as e:
        print(f"Error {tmpl.id}: {e}")
